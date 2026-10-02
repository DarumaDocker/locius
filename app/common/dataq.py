"""Exact table queries for the agent — counts, sums, averages, groups, percentiles, pivots, trends — over CSV / Excel
files and text logs, without running model-written code.

Why: in the 2026-10-02 round-5 test the model "read" a 40-row survey and reported 9 promoters / 19 detractors (NPS −25);
the file has 14 / 17 (NPS −7.5). Models can't count rows reliably, so every number that comes from a table goes through
this module. A query is a small JSON object (where / derive / group_by + agg / pivot / trend / sort / limit / save) applied
in a fixed order; expressions in derive use the safe calculator (app.common.calc)."""
from __future__ import annotations

import csv
import io
import math
import os
import re
import statistics
from datetime import datetime

from app.common import calc

MAX_ROWS = 500_000
MAX_BYTES = 60 * 1024 * 1024
SHOW_ROWS = 60


class DataError(Exception):
    pass


# ----------------------------------------------------------------------------------------------------------- loading
_NUM = re.compile(r"^\s*(?:S\$|US\$|HK\$|RM|[$€£¥])?\s*([-+]?(?:\d{1,3}(?:,\d{3})+|\d+)?(?:\.\d+)?(?:[eE][-+]?\d+)?)\s*(%)?\s*$")


def to_num(v):
    """'1,234.5' / 'S$ 80' / '12%' → float (12% → 12.0); anything else → None."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if not (isinstance(v, float) and math.isnan(v)) else None
    s = str(v).strip()
    if not s or len(s) > 40:
        return None
    m = _NUM.match(s)
    if not m or not re.search(r"\d", m.group(1) or ""):
        return None
    try:
        return float(m.group(1).replace(",", ""))
    except ValueError:
        return None


def _clean(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S").replace(" 00:00:00", "")
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return int(v)
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    n = to_num(s)
    if n is not None and not re.match(r"^0\d", s.lstrip("+-")):   # keep codes like 007 / phone numbers as text
        return int(n) if n.is_integer() and "%" not in s and abs(n) < 1e15 else n
    return s


def _read_text(path: str) -> str:
    raw = open(path, "rb").read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise DataError(f"file larger than {MAX_BYTES // 1024 // 1024} MB")
    for enc in ("utf-8-sig", "gb18030", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _dedupe(cols: list[str]) -> list[str]:
    out, seen = [], {}
    for i, c in enumerate(cols):
        c = str(c).strip() or f"col{i + 1}"
        if c in seen:
            seen[c] += 1
            c = f"{c}_{seen[c]}"
        else:
            seen[c] = 1
        out.append(c)
    return out


def load(path: str, sheet: str | None = None, pattern: str | None = None, header: bool = True) -> tuple[list[str], list[dict]]:
    """Read a table. CSV/TSV (delimiter sniffed), .xlsx (first sheet or `sheet`), or any text file with `pattern`
    (a regex with named groups: each matching line becomes a row)."""
    low = path.lower()
    if pattern:
        try:
            rx = re.compile(pattern)
        except re.error as e:
            raise DataError(f"bad pattern: {e}")
        if not rx.groupindex:
            raise DataError("pattern needs named groups, e.g. (?P<time>\\S+) (?P<status>\\d{3})")
        cols = list(rx.groupindex)
        rows, skipped = [], 0
        for line in _read_text(path).splitlines():
            m = rx.search(line)
            if not m:
                skipped += 1 if line.strip() else 0
                continue
            rows.append({c: _clean(m.group(c)) for c in cols})
            if len(rows) >= MAX_ROWS:
                break
        if not rows:
            raise DataError(f"pattern matched no lines ({skipped} non-empty lines did not match)")
        return cols, rows
    if low.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            raise DataError(f"cannot open workbook (damaged or not a real .xlsx?): {str(e)[:120]}")
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
        except KeyError:
            raise DataError(f"no sheet {sheet!r}; sheets: {wb.sheetnames}")
        it = ws.iter_rows(values_only=True)
        grid = []
        for r in it:
            if any(v not in (None, "") for v in r):
                grid.append(list(r))
            if len(grid) > MAX_ROWS:
                break
        wb.close()
    elif low.endswith((".csv", ".tsv", ".txt")):
        text = _read_text(path)
        sample = text[:20000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel_tab if low.endswith(".tsv") else csv.excel
        grid = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)][:MAX_ROWS + 1]
    else:
        raise DataError("supported: .csv .tsv .txt .xlsx (for logs or other text pass a regex `pattern`)")
    if not grid:
        raise DataError("the table is empty")
    width = max(len(r) for r in grid)
    if header:
        cols = _dedupe([str(c) if c is not None else "" for c in grid[0]] + [""] * (width - len(grid[0])))
        body = grid[1:]
    else:
        cols, body = [f"col{i + 1}" for i in range(width)], grid
    rows = [{c: _clean(r[i] if i < len(r) else None) for i, c in enumerate(cols)} for r in body]
    return cols, rows


# -------------------------------------------------------------------------------------------------------- operations
def alias(col: str) -> str:
    a = re.sub(r"\W+", "_", str(col)).strip("_") or "c"
    return ("c_" + a) if a[0].isdigit() else a


def _col(cols: list[str], name) -> str:
    n = str(name or "").strip()
    if n in cols:
        return n
    for c in cols:
        if alias(c) == n or c.lower() == n.lower() or alias(c).lower() == alias(n).lower():
            return c
    raise DataError(f"no column {name!r}; columns: {cols}")


def _cmp(a, op: str, b) -> bool:
    op = op.strip().lower()
    if op in ("empty", "is_empty"):
        return a in ("", None)
    if op in ("not_empty", "notempty"):
        return a not in ("", None)
    if op in ("in", "not_in"):
        vals = b if isinstance(b, list) else [b]
        hit = any(_eq(a, v) for v in vals)
        return hit if op == "in" else not hit
    if op == "between":
        if not isinstance(b, list) or len(b) != 2:
            raise DataError("between needs [low, high]")
        return _cmp(a, ">=", b[0]) and _cmp(a, "<=", b[1])
    if op in ("contains", "not_contains"):
        hit = str(b).lower() in str(a).lower()
        return hit if op == "contains" else not hit
    if op == "startswith":
        return str(a).startswith(str(b))
    if op == "regex":
        try:
            return re.search(str(b), str(a)) is not None
        except re.error as e:
            raise DataError(f"bad regex: {e}")
    if op in ("==", "=", "eq"):
        return _eq(a, b)
    if op in ("!=", "<>", "ne"):
        return not _eq(a, b)
    x, y = to_num(a), to_num(b)
    if x is None or y is None:
        x, y = str(a), str(b)      # dates / text compare as strings (ISO dates sort correctly)
    if op in (">", "gt"):
        return x > y
    if op in (">=", "ge"):
        return x >= y
    if op in ("<", "lt"):
        return x < y
    if op in ("<=", "le"):
        return x <= y
    raise DataError(f"unknown op {op!r}")


def _eq(a, b) -> bool:
    x, y = to_num(a), to_num(b)
    if x is not None and y is not None:
        return abs(x - y) < 1e-9
    return str(a).strip().lower() == str(b).strip().lower()


def _where(cols, rows, conds, any_of=False):
    """Conditions are ANDed (any_of: ORed). A value that names a column — or {"col": name} — compares the two columns
    row by row (e.g. 金额 > 预算)."""
    conds = conds if isinstance(conds, list) else [conds]
    cc = []
    for c in conds:
        if not isinstance(c, dict):
            raise DataError("each condition is an object {col, op, value}")
        v = c.get("value")
        other = None
        if isinstance(v, dict) and v.get("col"):
            other = _col(cols, v["col"])
        elif isinstance(v, str) and v in cols and c.get("col") != v:
            other = v
        cc.append((_col(cols, c.get("col")), str(c.get("op") or "=="), v, other))
    test = any if any_of else all
    return [r for r in rows if test(_cmp(r[c], op, r[o] if o else v) for c, op, v, o in cc)]


_DATE = re.compile(r"(\d{4})[-/.](\d{1,2})(?:[-/.](\d{1,2}))?(?:[T\s]+(\d{1,2}):(\d{2})(?::(\d{2}))?)?")


def _date_part(v, part: str):
    m = _DATE.search(str(v))
    if not m:
        return ""
    y, mo = int(m.group(1)), int(m.group(2))
    d = int(m.group(3) or 1)
    h = m.group(4)
    if part == "year":
        return y
    if part == "month":
        return f"{y:04d}-{mo:02d}"
    if part == "quarter":
        return f"{y}-Q{(mo - 1) // 3 + 1}"
    if part == "date":
        return f"{y:04d}-{mo:02d}-{d:02d}"
    if part == "hour":
        return int(h) if h is not None else ""
    if part == "weekday":
        try:
            return datetime(y, mo, d).strftime("%a")
        except ValueError:
            return ""
    if part == "week":
        try:
            iy, iw, _ = datetime(y, mo, d).isocalendar()
            return f"{iy}-W{iw:02d}"
        except ValueError:
            return ""
    raise DataError("date_part: year / quarter / month / week / date / weekday / hour")


def _derive(cols, rows, specs):
    for s in specs if isinstance(specs, list) else [specs]:
        name = str(s.get("as") or "").strip()
        if not name:
            raise DataError("derive needs \"as\" (the new column name)")
        if "expr" in s:
            names = {alias(c): c for c in cols}
            expr = str(s["expr"])
            for c in sorted(cols, key=len, reverse=True):   # 金额(新元) - 预算(新元) → 金额_新元 - 预算_新元
                if c != alias(c) and c in expr:
                    expr = expr.replace(c, alias(c))
            for r in rows:
                env = {}
                for a, c in names.items():
                    n = to_num(r[c])
                    env[a] = n if n is not None else 0.0
                try:
                    r[name] = calc.evaluate(expr, env)
                except calc.CalcError as e:
                    raise DataError(f"derive {name}: {e} (columns are referred to as {list(names)[:20]})")
                if isinstance(r[name], bool):
                    r[name] = int(r[name])
        elif "bins" in s:
            src = _col(cols, s.get("from"))
            edges = [float(x) for x in s["bins"]]
            labels = s.get("labels") or [f"{edges[i]:g}–{edges[i + 1]:g}" for i in range(len(edges) - 1)]
            if len(labels) != len(edges) - 1:
                raise DataError("bins: labels must be one fewer than edges")
            right = bool(s.get("right", False))       # default: [low, high)  ; right=true: (low, high]
            for r in rows:
                x, lab = to_num(r[src]), ""
                if x is not None:
                    for i in range(len(edges) - 1):
                        lo, hi = edges[i], edges[i + 1]
                        if (lo < x <= hi) if right else (lo <= x < hi):
                            lab = labels[i]
                            break
                r[name] = lab
        elif "map" in s:
            src, mp = _col(cols, s.get("from")), {str(k): v for k, v in dict(s["map"]).items()}
            for r in rows:
                r[name] = mp.get(str(r[src]), r[src] if s.get("keep", True) else s.get("default", ""))
        elif "date_part" in s:
            src = _col(cols, s.get("from"))
            for r in rows:
                r[name] = _date_part(r[src], str(s["date_part"]))
        elif "slice" in s:
            src, sl = _col(cols, s.get("from")), list(s["slice"]) + [None]
            for r in rows:
                r[name] = str(r[src])[sl[0]:sl[1]]
        elif "zscore" in s or "pct_change" in s or "rank" in s or "cumsum" in s:
            kind = next(k for k in ("zscore", "pct_change", "rank", "cumsum") if k in s)
            src = _col(cols, s[kind])
            by = [_col(cols, b) for b in (s.get("by") or [])]
            groups: dict = {}
            for r in rows:
                groups.setdefault(tuple(r[b] for b in by), []).append(r)
            for g in groups.values():
                xs = [to_num(r[src]) for r in g]
                if kind == "zscore":
                    vals = [x for x in xs if x is not None]
                    mu = statistics.fmean(vals) if vals else 0.0
                    sd = statistics.stdev(vals) if len(vals) > 1 else 0.0
                    for r, x in zip(g, xs):
                        r[name] = round((x - mu) / sd, 3) if (x is not None and sd) else ""
                elif kind == "pct_change":
                    prev = None
                    for r, x in zip(g, xs):
                        r[name] = round((x - prev) / abs(prev) * 100, 2) if (x is not None and prev) else ""
                        prev = x if x is not None else prev
                elif kind == "cumsum":
                    tot = 0.0
                    for r, x in zip(g, xs):
                        tot += x or 0.0
                        r[name] = tot
                else:
                    order = sorted(range(len(g)), key=lambda i: -(xs[i] if xs[i] is not None else -math.inf))
                    for k, i in enumerate(order, 1):
                        g[i][name] = k
        else:
            raise DataError("derive needs one of: expr, bins, map, date_part, slice, zscore, pct_change, rank, cumsum")
        if name not in cols:
            cols.append(name)
    return cols, rows


def _pct(vals: list[float], p: float) -> float:
    s = sorted(vals)
    if not s:
        return float("nan")
    k = (len(s) - 1) * p / 100
    f, c = math.floor(k), math.ceil(k)
    return s[f] if f == c else s[f] + (s[c] - s[f]) * (k - f)


def _agg(vals_raw: list, fn: str, total=None):
    fn = fn.strip().lower()
    nums = [x for x in (to_num(v) for v in vals_raw) if x is not None]
    if fn == "count":
        return len(vals_raw)
    if fn in ("count_nonempty", "nonempty"):
        return sum(1 for v in vals_raw if v not in ("", None))
    if fn in ("count_distinct", "nunique", "distinct"):
        return len({str(v) for v in vals_raw if v not in ("", None)})
    if fn == "mode":
        vs = [str(v) for v in vals_raw if v not in ("", None)]
        return statistics.multimode(vs)[0] if vs else ""
    if fn in ("first", "last"):
        return (vals_raw[0] if fn == "first" else vals_raw[-1]) if vals_raw else ""
    if fn in ("list", "values"):
        return ", ".join(dict.fromkeys(str(v) for v in vals_raw if v not in ("", None)))[:300]
    if not nums:
        return ""
    if fn == "sum":
        return sum(nums)
    if fn in ("mean", "avg", "average"):
        return statistics.fmean(nums)
    if fn == "median":
        return statistics.median(nums)
    if fn == "min":
        return min(nums)
    if fn == "max":
        return max(nums)
    if fn in ("std", "stdev"):
        return statistics.stdev(nums) if len(nums) > 1 else 0.0
    if fn in ("pstd", "pstdev"):
        return statistics.pstdev(nums)
    if fn in ("var", "variance"):
        return statistics.variance(nums) if len(nums) > 1 else 0.0
    if fn == "range":
        return max(nums) - min(nums)
    m = re.fullmatch(r"p(\d{1,2}(?:\.\d+)?)", fn)
    if m:
        return _pct(nums, float(m.group(1)))
    if fn in ("share", "sum_share"):
        return sum(nums) / total * 100 if total else ""
    raise DataError(f"unknown agg fn {fn!r} (count, count_distinct, sum, mean, median, min, max, std, var, range, "
                    "p50/p90/p95/p99, share, count_share, mode, first, last, list)")


def _group(cols, rows, by, aggs):
    by = [_col(cols, b) for b in (by if isinstance(by, list) else [by] if by else [])]
    aggs = aggs if isinstance(aggs, list) else [aggs]
    if not aggs:
        aggs = [{"fn": "count"}]
    plan = []
    for a in aggs:
        fn = str(a.get("fn") or "count")
        c = _col(cols, a["col"]) if a.get("col") not in (None, "", "*") else None
        if c is None and fn.lower() not in ("count", "count_share"):
            raise DataError(f"agg {fn} needs \"col\"")
        plan.append((c, fn, str(a.get("as") or (f"{fn}({c})" if c else fn))))
    groups: dict = {}
    for r in rows:
        groups.setdefault(tuple(r[b] for b in by), []).append(r)
    totals = {c: sum(x for x in (to_num(r[c]) for r in rows) if x is not None) for c, fn, _ in plan
              if c and fn.lower() in ("share", "sum_share")}
    out_cols = by + [p[2] for p in plan]
    out = []
    for key, g in groups.items():
        o = dict(zip(by, key))
        for c, fn, name in plan:
            if fn.lower() == "count_share":
                o[name] = len(g) / len(rows) * 100 if rows else ""
            else:
                o[name] = _agg([r[c] for r in g] if c else g, fn, totals.get(c))
        out.append(o)
    return out_cols, out


def _pivot(cols, rows, spec):
    rc, cc = _col(cols, spec.get("rows")), _col(cols, spec.get("cols"))
    vc = _col(cols, spec["value"]) if spec.get("value") else None
    fn = str(spec.get("fn") or ("sum" if vc else "count"))
    rkeys = list(dict.fromkeys(r[rc] for r in rows))
    ckeys = list(dict.fromkeys(r[cc] for r in rows))
    if spec.get("sort", True):
        rkeys, ckeys = sorted(rkeys, key=str), sorted(ckeys, key=str)
    cell: dict = {}
    for r in rows:
        cell.setdefault((r[rc], r[cc]), []).append(r[vc] if vc else 1)
    totals = bool(spec.get("totals", True))
    out_cols = [rc] + [str(k) for k in ckeys] + (["Total"] if totals else [])
    out = []
    for rk in rkeys:
        o = {rc: rk}
        allv = []
        for ck in ckeys:
            vs = cell.get((rk, ck), [])
            allv += vs
            o[str(ck)] = _agg(vs, fn) if vs else 0 if fn in ("sum", "count") else ""
        if totals:
            o["Total"] = _agg(allv, fn)
        out.append(o)
    if totals:
        t = {rc: "Total"}
        for ck in ckeys:
            t[str(ck)] = _agg([v for (a, b), vs in cell.items() if b == ck for v in vs], fn)
        t["Total"] = _agg([v for vs in cell.values() for v in vs], fn)
        out.append(t)
    return out_cols, out


def _trend(cols, rows, spec):
    """Least-squares line y = a + b·x per group; x defaults to the row order 1..n (e.g. months in order)."""
    yc = _col(cols, spec["y"])
    xc = _col(cols, spec["x"]) if spec.get("x") else None
    by = [_col(cols, b) for b in (spec.get("by") or [])]
    ahead = max(0, min(int(spec.get("ahead", 1) or 0), 24))
    groups: dict = {}
    for r in rows:
        groups.setdefault(tuple(r[b] for b in by), []).append(r)
    out_cols = by + ["n", "slope", "intercept", "r2", "last_y"] + [f"forecast_{i}" for i in range(1, ahead + 1)]
    out = []
    for key, g in groups.items():
        pts = []
        for i, r in enumerate(g, 1):
            x = to_num(r[xc]) if xc else float(i)
            y = to_num(r[yc])
            if x is not None and y is not None:
                pts.append((x, y))
        o = dict(zip(by, key))
        o["n"] = len(pts)
        if len(pts) < 2:
            out.append(o)
            continue
        mx, my = statistics.fmean(p[0] for p in pts), statistics.fmean(p[1] for p in pts)
        sxx = sum((p[0] - mx) ** 2 for p in pts)
        b = sum((p[0] - mx) * (p[1] - my) for p in pts) / sxx if sxx else 0.0
        a = my - b * mx
        ss_tot = sum((p[1] - my) ** 2 for p in pts)
        ss_res = sum((p[1] - (a + b * p[0])) ** 2 for p in pts)
        o.update(slope=b, intercept=a, r2=(1 - ss_res / ss_tot) if ss_tot else 1.0, last_y=pts[-1][1])
        step = (pts[-1][0] - pts[0][0]) / (len(pts) - 1) if len(pts) > 1 else 1.0
        for i in range(1, ahead + 1):
            o[f"forecast_{i}"] = a + b * (pts[-1][0] + step * i)
        out.append(o)
    return out_cols, out


def _sort(cols, rows, spec):
    for s in reversed(spec if isinstance(spec, list) else [spec]):
        if isinstance(s, str):
            s = {"col": s.lstrip("-"), "desc": s.startswith("-")}
        c = _col(cols, s["col"])
        desc = bool(s.get("desc"))
        num = all(to_num(r[c]) is not None for r in rows if r[c] not in ("", None))
        if num:
            rows.sort(key=lambda r: (r[c] in ("", None), (-1 if desc else 1) * (to_num(r[c]) or 0.0)))
        else:
            rows.sort(key=lambda r: str(r[c]), reverse=desc)
    return rows


def run(path: str, q: dict) -> tuple[list[str], list[dict], dict]:
    """Apply one query. Returns (columns, rows, info)."""
    cols, rows = load(path, q.get("sheet"), q.get("pattern"), q.get("header", True))
    info = {"source_rows": len(rows), "columns": list(cols)}
    if q.get("derive"):     # first, so where / group_by can use derived columns (z-scores see the whole table)
        cols, rows = _derive(list(cols), rows, q["derive"])
    if q.get("where"):
        rows = _where(cols, rows, q["where"])
    if q.get("where_any"):
        rows = _where(cols, rows, q["where_any"], any_of=True)
    info["matched_rows"] = len(rows)
    if q.get("pivot"):
        cols, rows = _pivot(cols, rows, q["pivot"])
    elif q.get("trend"):
        cols, rows = _trend(cols, rows, q["trend"])
    elif q.get("group_by") or q.get("agg"):
        cols, rows = _group(cols, rows, q.get("group_by") or [], q.get("agg") or [{"fn": "count"}])
    if q.get("having"):     # filter the grouped / pivoted result
        rows = _where(cols, rows, q["having"])
    if q.get("sort"):
        rows = _sort(cols, rows, q["sort"])
    if q.get("select"):
        cols = [_col(cols, c) for c in q["select"]]
    if q.get("totals") and not q.get("pivot"):
        t = {}
        for c in cols:
            nums = [to_num(r.get(c)) for r in rows]
            t[c] = sum(x for x in nums if x is not None) if nums and all(x is not None for x in nums) else ""
        t[cols[0]] = "Total"
        rows = rows + [t]
    if q.get("limit"):
        rows = rows[:max(1, int(q["limit"]))]
    return cols, rows, info


def describe(path: str, sheet=None, pattern=None, header=True) -> str:
    cols, rows = load(path, sheet, pattern, header)
    lines = [f"{len(rows)} rows × {len(cols)} columns. Refer to columns by name (or by the alias shown)."]
    for c in cols:
        vals = [r[c] for r in rows]
        nonempty = [v for v in vals if v not in ("", None)]
        nums = [x for x in (to_num(v) for v in nonempty) if x is not None]
        al = alias(c)
        head = f"- {c}" + (f" (alias {al})" if al != c else "")
        if nonempty and len(nums) == len(nonempty):
            lines.append(f"{head}: number, {len(nums)} values, min {fmt(min(nums))}, max {fmt(max(nums))}, "
                         f"mean {fmt(statistics.fmean(nums))}, sum {fmt(sum(nums))}")
        else:
            distinct: dict = {}
            for v in nonempty:
                distinct[str(v)] = distinct.get(str(v), 0) + 1
            top = sorted(distinct.items(), key=lambda kv: -kv[1])[:6]
            lines.append(f"{head}: text, {len(nonempty)} non-empty, {len(distinct)} distinct; top: "
                         + ", ".join(f"{k[:30]} ({n})" for k, n in top))
    lines.append("First rows:")
    lines.append(table(cols, rows[:5]))
    return "\n".join(lines)


def fmt(v) -> str:
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        if v.is_integer() and abs(v) < 1e15:
            return f"{int(v):,}"
        return f"{v:,.4f}".rstrip("0").rstrip(".") if abs(v) >= 0.0001 else f"{v:.4g}"
    if isinstance(v, int) and not isinstance(v, bool):
        return f"{v:,}" if abs(v) >= 10000 else str(v)
    return str(v).replace("|", "/").replace("\n", " ")


def table(cols: list[str], rows: list[dict], limit: int = SHOW_ROWS) -> str:
    if not rows:
        return "(no rows)"
    out = ["| " + " | ".join(str(c) for c in cols) + " |", "|" + "---|" * len(cols)]
    for r in rows[:limit]:
        out.append("| " + " | ".join(fmt(r.get(c, "")) for c in cols) + " |")
    if len(rows) > limit:
        out.append(f"… {len(rows) - limit} more rows (use save to get them all, or group/limit)")
    return "\n".join(out)


def save(cols: list[str], rows: list[dict], path: str) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    if path.lower().endswith(".xlsx"):
        from app.common import xlsx
        xlsx.make(path, [{"name": "Data", "columns": cols, "rows": [[r.get(c, "") for c in cols] for r in rows]}])
        return
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for r in rows:
            w.writerow([r.get(c, "") for c in cols])
