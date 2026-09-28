"""0.2.7: step budget, Excel export, blocked-site handling, Google Calendar connector."""
import json
import os
import tempfile

import httpx
import pytest

from app.common import xlsx
from app.runtime.agent import (BUDGET_MARK, BUDGET_MARK_PAGES, FINISH_TOOLS, LOCAL_TOOLS, Runtime, site_blocked,
                               site_brand)
from app.sentinel import gcal
from app.sentinel.policy import ALLOW, ASK, DENY, decide
from app.sentinel.store import Store


# ---------------------------------------------------------------- step budget
def _tools(*names):
    return [{"type": "function", "function": {"name": n}} for n in names]


def _nav_turn(i):
    return {"role": "assistant", "content": "", "tool_calls": [{"id": f"c{i}", "type": "function",
                                                                "function": {"name": "browser_navigate", "arguments": "{}"}}]}


def test_budget_note_once_and_finish_tools_only():
    tr = [{"role": "system", "content": ""}, {"role": "user", "content": "research X and make a PDF"}]
    all_tools = _tools("browser_navigate", "browser_snapshot", "make_pdf", "send_file", "files_write", "gmail_send")
    out = Runtime._budget(tr, all_tools, 20)
    assert out == all_tools and len(tr) == 2
    out = Runtime._budget(tr, all_tools, 5)
    assert BUDGET_MARK in tr[-1]["content"] and len(out) == len(all_tools)
    n = len(tr)
    Runtime._budget(tr, all_tools, 4)
    assert len(tr) == n                      # said only once
    out = Runtime._budget(tr, all_tools, 2)
    names = {t["function"]["name"] for t in out}
    assert names == {"make_pdf", "send_file", "files_write", "gmail_send"}
    assert "browser_navigate" not in FINISH_TOOLS


def test_budget_research_nudge_after_many_pages():
    tr = [{"role": "system", "content": ""}, {"role": "user", "content": "news brief"}]
    tr += [_nav_turn(i) for i in range(10)]
    Runtime._budget(tr, _tools("browser_navigate"), 25)
    assert BUDGET_MARK_PAGES in tr[-1]["content"] and "10" in tr[-1]["content"]
    n = len(tr)
    tr.append(_nav_turn(11))
    Runtime._budget(tr, _tools("browser_navigate"), 24)
    assert len(tr) == n + 1                  # nudged once


def test_make_xlsx_tool_is_declared():
    assert any(t["function"]["name"] == "make_xlsx" for t in LOCAL_TOOLS)


# ---------------------------------------------------------------- xlsx
def test_xlsx_roundtrip_numbers_and_safe_formulas(tmp_path):
    p = str(tmp_path / "t.xlsx")
    info = xlsx.make(p, [{"name": "航班 Flights", "columns": ["航空公司", "价格 SGD", "占比", "备注"],
                          "rows": [["Scoot", "600", "12.5%", "=HYPERLINK(\"http://evil\",\"x\")"],
                                   ["ANA", "1,085", "30%", "=SUM(B2:B3)"],
                                   ["ZIPAIR", 805.5, None, "-"]]}], formulas=True)
    assert info == {"sheets": ["航班 Flights"], "rows": 3}
    from openpyxl import load_workbook
    ws = load_workbook(p).active
    assert ws["B2"].value == 600 and ws["B3"].value == 1085 and ws["B4"].value == 805.5
    assert abs(ws["C2"].value - 0.125) < 1e-9 and ws["C2"].number_format == "0.0%"
    assert ws["D2"].data_type == "s"                      # HYPERLINK never becomes a live formula
    assert ws["D3"].value == "=SUM(B2:B3)" and ws["D3"].data_type == "f"
    assert ws.freeze_panes == "A2" and ws["A1"].font.bold
    txt = xlsx.read_text(p)
    assert "Scoot" in txt and "航班 Flights" in txt


def test_xlsx_formulas_off_by_default(tmp_path):
    p = str(tmp_path / "t.xlsx")
    xlsx.make(p, [{"columns": ["a"], "rows": [["=1+1"], ["+cmd|' /C calc'!A0"]]}])
    from openpyxl import load_workbook
    ws = load_workbook(p).active
    assert ws["A2"].data_type == "s" and ws["A3"].data_type == "s"


def test_xlsx_from_markdown_and_csv():
    rows = xlsx.rows_from_markdown("intro\n\n| 名称 | 价格 |\n|---|---:|\n| **A** | 1 |\n| B | 2 |\n\nafter")
    assert rows == [["名称", "价格"], ["A", "1"], ["B", "2"]]
    assert xlsx.rows_from_csv('a,b\n"x, y",2\n') == [["a", "b"], ["x, y", "2"]]


# ---------------------------------------------------------------- blocked sites
@pytest.mark.parametrize("url,brand", [
    ("https://www.opentable.com/s?x=1", "opentable"), ("https://m.opentable.sg/r/abc", "opentable"),
    ("opentable.sg", "opentable"), ("https://www.bbc.co.uk/news", "bbc"), ("https://shop.test:8099/a", "shop"),
    ("https://www.google.com.sg/maps", "google"),
])
def test_site_brand(url, brand):
    assert site_brand(url) == brand


def test_site_blocked_until_takeover():
    tr = [{"role": "tool", "content": "[SITE BLOCKED site=opentable] wall\n<untrusted_content>…"}]
    assert site_blocked(tr, "https://www.opentable.sg/singapore") == "opentable"
    assert site_blocked(tr, "https://www.chope.co/") == ""
    tr.append({"role": "tool", "content": "用户已完成接管并交还浏览器控制权。"})
    assert site_blocked(tr, "https://www.opentable.com/") == ""


# ---------------------------------------------------------------- Google Calendar client (fake Google API)
class FakeGoogle:
    def __init__(self):
        self.calls = []
        self.events = [
            {"id": "ev1", "summary": "Standup", "start": {"dateTime": "2026-10-05T09:30:00+08:00"},
             "end": {"dateTime": "2026-10-05T10:00:00+08:00"}, "attendees": [{"email": "a@x.com", "responseStatus": "accepted"}]},
            {"id": "ev2", "summary": "Holiday", "start": {"date": "2026-10-06"}, "end": {"date": "2026-10-07"}},
        ]

    def __call__(self, req: httpx.Request):
        self.calls.append((req.method, req.url.path, dict(req.url.params), req.content.decode() if req.content else ""))
        path = req.url.path
        if path.endswith("/token"):
            return httpx.Response(200, json={"access_token": "at-1", "expires_in": 3600})
        if req.headers.get("authorization") != "Bearer at-1":
            return httpx.Response(401, json={"error": {"message": "bad token"}})
        if path.endswith("/freeBusy"):
            return httpx.Response(200, json={"calendars": {"primary": {"busy": [
                {"start": "2026-10-05T01:30:00Z", "end": "2026-10-05T02:00:00Z"},
                {"start": "2026-10-05T04:00:00Z", "end": "2026-10-05T06:00:00Z"}]}}})
        if req.method == "GET" and path.endswith("/events"):
            return httpx.Response(200, json={"items": self.events})
        if req.method == "POST" and path.endswith("/events"):
            b = json.loads(req.content)
            return httpx.Response(200, json={"id": "new1", **b, "htmlLink": "https://calendar.google.com/x"})
        if req.method == "GET" and "/events/" in path:
            return httpx.Response(200, json=self.events[0])
        if req.method == "PATCH":
            return httpx.Response(200, json={**self.events[0], **json.loads(req.content)})
        if req.method == "DELETE":
            return httpx.Response(204)
        return httpx.Response(404, json={"error": {"message": "nope"}})


@pytest.fixture
def cal():
    fake = FakeGoogle()
    g = gcal.GCal("cid", "sec", "rt", "Asia/Singapore")
    g.c = httpx.Client(transport=httpx.MockTransport(fake))
    g.fake = fake
    return g


def test_gcal_list_events(cal):
    r = cal.list_events("2026-10-05", "2026-10-08")
    assert [e["title"] for e in r["events"]] == ["Standup", "Holiday"]
    assert r["events"][0]["start"] == "2026-10-05 09:30" and r["events"][1]["all_day"]
    assert r["events"][0]["attendees"] == ["a@x.com (accepted)"]
    m, path, params, _ = cal.fake.calls[-1]
    assert params["singleEvents"] == "true" and params["timeMin"].startswith("2026-10-05T00:00:00+08:00")


def test_gcal_free_slots(cal):
    r = cal.free_slots("2026-10-05 09:00", "2026-10-05 18:00", 60)
    # busy 09:30–10:00 and 12:00–14:00 (SGT) → free 10:00–12:00 and 14:00–18:00 (09:00–09:30 is too short)
    assert r["free"] == ["2026-10-05 Mon 10:00–12:00", "2026-10-05 Mon 14:00–18:00"]


def test_gcal_create_with_invites(cal):
    r = cal.create_event({"title": "🍽 Kuriya", "start": "2026-10-03 19:00", "location": "Orchard",
                          "attendees": "eva@example.com, bad-address", "reminder_minutes": 120})
    assert r["created"] and r["invites_sent_to"] == ["eva@example.com"]
    m, path, params, body = cal.fake.calls[-1]
    b = json.loads(body)
    assert params["sendUpdates"] == "all" and b["start"]["dateTime"] == "2026-10-03T19:00:00+08:00"
    assert b["end"]["dateTime"] == "2026-10-03T20:00:00+08:00" and b["reminders"]["overrides"][0]["minutes"] == 120


def test_gcal_update_keeps_duration_and_delete(cal):
    r = cal.update_event({"event_id": "ev1", "start": "2026-10-05 11:00"})
    body = json.loads(cal.fake.calls[-1][3])
    assert body["start"]["dateTime"] == "2026-10-05T11:00:00+08:00" and body["end"]["dateTime"] == "2026-10-05T11:30:00+08:00"
    assert cal.delete_event({"event_id": "ev1"})["deleted"]
    with pytest.raises(gcal.GCalError):
        cal.delete_event({"event_id": "../../x"})


def test_gcal_bad_time(cal):
    with pytest.raises(gcal.GCalError):
        cal.create_event({"title": "x", "start": "next friday"})
    with pytest.raises(gcal.GCalError):
        cal.create_event({"title": "x", "start": "2026-10-03 19:00", "end": "2026-10-03 18:00"})


def test_gcal_auth_url():
    u = gcal.auth_url("1-a.apps.googleusercontent.com", "https://box/sentinel/api/connections/calendar/callback", "st")
    assert "access_type=offline" in u and "calendar.events" in u and "state=st" in u


# ---------------------------------------------------------------- calendar policy
@pytest.fixture
def store():
    return Store(tempfile.mkdtemp())


def test_policy_calendar(store):
    assert decide(store, "calendar_list_events", {}, "t1").decision == DENY          # not connected
    store.put_secret("calendar", {"client_id": "c", "client_secret": "s", "refresh_token": "r"})
    store.save_connection("calendar", {"email": "me@x.com"}, enabled=True)
    assert decide(store, "calendar_list_events", {}, "t1").decision == ALLOW
    assert decide(store, "calendar_create_event", {"title": "x", "start": "2026-10-03 19:00"}, "t1").decision == ALLOW
    assert decide(store, "calendar_create_event", {"title": "x", "start": "2026-10-03 19:00",
                                                    "attendees": ["a@b.com"]}, "t1").decision == ASK
    assert decide(store, "calendar_update_event", {"event_id": "e"}, "t1").decision == ASK
    assert decide(store, "calendar_delete_event", {"event_id": "e"}, "t1").decision == ASK
    store.save_connection("calendar", permissions={"write": False})
    assert decide(store, "calendar_create_event", {"title": "x", "start": "2026-10-03"}, "t1").decision == DENY
