"""0.2.7 end-to-end on the local stack: step budget keeps the last steps for the PDF, bot walls stop retries,
Excel export, Google Calendar OAuth connect + free slots + invite (approval)."""
import os
import sys
import time

import httpx

B = "http://127.0.0.1:8080"
H = {"X-Persona-UI": "1"}
T = "/tmp/claude-0/persona-test"
WS = T + "/workspace"
c = httpx.Client(timeout=90, trust_env=False)
fails = []


def check(name, cond, info=""):
    print(("PASS " if cond else "FAIL ") + name, "" if cond else str(info)[:1500])
    if not cond:
        fails.append(name)


def wait(tid, timeout=120, until=("COMPLETED", "FAILED", "CANCELLED")):
    t0 = time.time()
    while time.time() - t0 < timeout:
        t = c.get(f"{B}/api/tasks/{tid}").json()
        if t["status"] in until:
            return t
        time.sleep(0.5)
    return t


def chat(msg):
    return c.post(B + "/api/chat", json={"message": msg}, headers=H).json()


# ------------------------------------------------------------------ 1. step budget
old = c.get(B + "/api/settings").json()
c.put(B + "/api/settings", json={"max_steps": 8}, headers=H)
r = chat("BUDGETPDF 调研日本餐厅并生成 PDF 简报发给我")
t = wait(r["task_id"])
navs = [e for e in t["events"] if e["type"] == "tool_call" and e["data"]["name"] == "browser_navigate"]
sent = [e for e in t["events"] if e["type"] == "tool_call" and e["data"]["name"] == "send_file"]
check("budget: task completed (not FAILED at the limit)", t["status"] == "COMPLETED", (t["status"], t.get("error"), t.get("result")))
check("budget: stopped browsing and delivered the PDF", len(sent) == 1 and os.path.isfile(WS + "/reports/brief.pdf") and 2 <= len(navs) <= 4,
      (len(navs), len(sent)))
conv = c.get(f"{B}/api/conversations/{r['conversation_id']}").json()
check("budget: PDF card in the chat", any('"type": "file"' in m["content"] and "brief.pdf" in m["content"] for m in conv["messages"]))
c.put(B + "/api/settings", json={"max_steps": old.get("max_steps", 40)}, headers=H)

# ------------------------------------------------------------------ 2. blocked site
r = chat("BLOCKSITE 在 OpenTable 查餐厅")
t = wait(r["task_id"])
res = t.get("result") or ""
parts = res.split("=====")
check("blocked: wall detected on first visit", len(parts) == 3 and "[SITE BLOCKED site=opentable]" in parts[0], res[:600])
check("blocked: second URL on the same site not even opened", "ERROR: [SITE BLOCKED site=opentable]" in parts[1], res[:900])
check("blocked: other sites still work", "Kuriya" in parts[2] and "SITE BLOCKED" not in parts[2], parts[2:] and parts[2][:300])
ev = [e for e in t["events"] if e["type"] == "site_blocked"]
check("blocked: site_blocked event recorded", len(ev) == 1 and ev[0]["data"]["kind"] in ("akamai", "http-403"), ev)
bnav = [e for e in t["events"] if e["type"] == "tool_result" and e["data"]["name"] == "browser_navigate"]
check("blocked: 3 navigate results", len(bnav) == 3, len(bnav))

# ------------------------------------------------------------------ 3. Excel
r = chat("XLSXOUT 把航班做成 Excel 发给我")
t = wait(r["task_id"])
p = WS + "/reports/flights.xlsx"
check("xlsx: task completed", t["status"] == "COMPLETED", (t["status"], t.get("result")))
ok = os.path.isfile(p)
if ok:
    from openpyxl import load_workbook
    ws = load_workbook(p).active
    ok = ws.title == "直飞航班" and ws["C4"].value == 1085 and ws["A1"].value == "航空公司"
check("xlsx: workbook has the rows with numbers as numbers", ok)
raw = c.get(B + "/api/files/raw", params={"path": "reports/flights.xlsx", "download": "1"})
check("xlsx: downloadable from the chat", raw.status_code == 200 and raw.content[:2] == b"PK", raw.status_code)
fr = c.post(B + "/api/chat", json={"message": "hi"}, headers=H)  # keep the stack warm

# ------------------------------------------------------------------ 4. Google Calendar
c.post("http://127.0.0.1:8094/_g/reset")
cat = c.get(B + "/sentinel/api/connections", headers=H).json()
cal = next(x for x in cat["connections"] if x["name"] == "calendar")
check("calendar: listed, not connected yet", not cal["has_credential"] and not cal["enabled"], cal)
bad = c.post(B + "/sentinel/api/connections/calendar/start", json={"client_id": "nope", "client_secret": "x"}, headers=H)
check("calendar: rejects a malformed client id", bad.status_code == 400, bad.text)
st = c.post(B + "/sentinel/api/connections/calendar/start", headers=H,
            json={"client_id": "1234567890-abcdef.apps.googleusercontent.com", "client_secret": "GOCSPX-test-secret-123"}).json()
check("calendar: consent URL points at Google auth with offline access", st["auth_url"].startswith("http://127.0.0.1:8094/g/auth?")
      and "access_type=offline" in st["auth_url"] and st["redirect_uri"].endswith("/sentinel/api/connections/calendar/callback"), st)
# forged callback with the wrong state is refused
forged = c.get(B + "/sentinel/api/connections/calendar/callback", params={"code": "x", "state": "evil"})
check("calendar: callback with a wrong state is refused", "state mismatch" in forged.text and forged.status_code == 200, forged.text[:200])
# the user's browser follows Google's redirect back to OMuse
g = c.get(st["auth_url"], follow_redirects=False)
loc = g.headers["location"]
cb = c.get(loc.replace(st["redirect_uri"].split("/sentinel/")[0], B))
check("calendar: callback stores the token and shows success", "lucas@example.com" in cb.text and "✓" in cb.text, cb.text[:300])
replay = c.get(loc.replace(st["redirect_uri"].split("/sentinel/")[0], B))
check("calendar: the same callback can't be replayed", "state mismatch" in replay.text, replay.text[:200])
cal = next(x for x in c.get(B + "/sentinel/api/connections", headers=H).json()["connections"] if x["name"] == "calendar")
check("calendar: connected, time zone read from Google", cal["has_credential"] and cal["enabled"] and cal["config"]["time_zone"] == "Asia/Singapore"
      and "refresh_token" not in str(cal), cal)
tst = c.post(B + "/sentinel/api/connections/calendar/test", json={}, headers=H).json()
check("calendar: test button works", tst.get("ok") and tst.get("upcoming") == 1, tst)

r = chat("CALBOOK 周六晚上帮我把订位放进日历，并邀请 Eva")
t = wait(r["task_id"], until=("WAITING_APPROVAL", "COMPLETED", "FAILED"))
check("calendar: invite needs approval", t["status"] == "WAITING_APPROVAL", (t["status"], t.get("result"), t.get("error")))
fs = [e for e in t["events"] if e["type"] == "tool_result" and e["data"]["name"] == "calendar_free_slots"]
check("calendar: free slots skip the 18:00 board meeting", fs and "2026-10-03 Sat 18:45–22:00" in fs[0]["data"]["preview"], fs and fs[0]["data"]["preview"])
aps = c.get(B + "/sentinel/api/approvals", params={"status": "pending"}, headers=H).json()
ap = next(a for a in aps["approvals"] if a["task_id"] == r["task_id"])
fields = dict(tuple(x) for x in ap["summary"]["fields"])
check("calendar: approval shows the invitee and time", "eva@example.com" in str(fields) and fields.get("开始 Start") == "2026-10-03 19:00"
      and "start" in ap["summary"]["editable"], ap["summary"])
res = c.post(f"{B}/sentinel/api/approvals/{ap['id']}/resolve", json={"decision": "approve", "scope": "ONCE",
                                                                    "args": {"start": "2026-10-03 19:30"}}, headers=H)
t = wait(r["task_id"])
gs = c.get("http://127.0.0.1:8094/_g").json()
made = [x for x in gs["log"] if "create" in x]
check("calendar: event created with the edited time and invites sent", t["status"] == "COMPLETED" and made
      and made[0]["create"]["start"]["dateTime"] == "2026-10-03T19:30:00+08:00" and made[0]["sendUpdates"] == "all"
      and made[0]["create"]["attendees"] == [{"email": "eva@example.com"}], (t["status"], made, t.get("result")))
tok_log = [x for x in gs["log"] if "token" in x]
check("calendar: access token reused (no refresh per call)", sum(1 for x in tok_log if x["token"].get("grant_type") == "refresh_token") <= 1, tok_log)
d = c.delete(B + "/sentinel/api/connections/calendar/credential", headers=H)
cal = next(x for x in c.get(B + "/sentinel/api/connections", headers=H).json()["connections"] if x["name"] == "calendar")
check("calendar: disconnect removes the credential", d.status_code == 200 and not cal["has_credential"] and not cal["enabled"], cal)

print("\n" + ("ALL PASS" if not fails else f"{len(fails)} FAILED: {fails}"))
sys.exit(1 if fails else 0)
